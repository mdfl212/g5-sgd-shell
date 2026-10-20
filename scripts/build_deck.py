#!/usr/bin/env python3
"""Build an editable G5-style PowerPoint from a JSON outline.

Prototype renderer: supports text, cards, tags, numbered lists, tables, and a
simple three-node flowchart. Figure placement and advanced flowcharts are not
yet implemented. Content accuracy remains the responsibility of the source
and review workflow.
"""
from __future__ import annotations
import argparse, json, sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
THEME_PATH = ROOT / "theme.json"
EXPECTED_STEPS = ["cover", "case", "questions", "answers", "visuals", "takeaways", "references", "thank-you"]
DEFAULT_ROSTER = "BISTER, DEIPARINE, DATO, KANASHIRO, MONTEROSO, & TIBURON"


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def add_text(slide, text, x, y, w, h, size=14, color="202124", bold=False,
             font="Cambria", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margin=0.06):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.025); tf.margin_bottom = Inches(0.025)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run(); run.text = str(text)
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line)
    return shape


def validate_outline(data, theme):
    errors = []
    if data.get("schemaVersion") != "1.0": errors.append("schemaVersion must be '1.0'.")
    meta = data.get("metadata", {})
    if meta.get("courseCode") not in theme["subjects"]: errors.append("metadata.courseCode must be MD201–MD207.")
    if not meta.get("title"): errors.append("metadata.title is required.")
    slides = data.get("slides")
    if not isinstance(slides, list) or not slides: return errors + ["slides must be a non-empty list."]
    allowed = {x["id"] for x in theme["lockedSteps"]}
    steps = []
    for i, slide in enumerate(slides, 1):
        if slide.get("step") not in allowed: errors.append(f"Slide {i}: unknown step {slide.get('step')!r}.")
        else: steps.append(slide["step"])
        if not slide.get("title"): errors.append(f"Slide {i}: title is required.")
        if not isinstance(slide.get("content", []), list): errors.append(f"Slide {i}: content must be a list.")
    first_occurrences = []
    for step in steps:
        if step not in first_occurrences: first_occurrences.append(step)
    if first_occurrences != EXPECTED_STEPS:
        errors.append("First occurrences of steps must be exactly: " + " -> ".join(EXPECTED_STEPS))
    for required in EXPECTED_STEPS:
        if required not in steps: errors.append(f"Missing required step: {required}.")
    return errors


def add_shell(slide, meta, theme, slide_no, slide_count):
    course = meta["courseCode"]
    palette = theme["subjects"][course]
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb("FFFFFF")
    add_rect(slide, 0, 0, 3.4, 0.11, palette["banner"], radius=False)
    add_rect(slide, 3.4, 0, 3.3, 0.11, palette["subhead"], radius=False)
    add_rect(slide, 6.7, 0, 3.3, 0.11, palette["accent"], radius=False)
    logo_path = ROOT / theme["branding"]["logo"]
    if logo_path.exists():
        slide.shapes.add_picture(str(logo_path), Inches(0.28), Inches(0.20), height=Inches(0.47))
    add_text(slide, theme["branding"]["university"], 0.83, 0.18, 4.0, 0.18, size=8, color=palette["banner"], bold=True)
    add_text(slide, theme["branding"]["school"], 0.83, 0.36, 3.5, 0.17, size=8, color="333333", bold=True)
    add_text(slide, theme["branding"]["campusLine"], 0.83, 0.53, 5.7, 0.15, size=5.7, color="555555")
    add_text(slide, course.replace("MD", "MD "), 8.38, 0.25, 1.25, 0.30, size=12, color=palette["banner"], bold=True, align=PP_ALIGN.RIGHT)
    add_rect(slide, 0, 5.48, 3.4, 0.07, palette["banner"], radius=False)
    add_rect(slide, 3.4, 5.48, 3.3, 0.07, palette["subhead"], radius=False)
    add_rect(slide, 6.7, 5.48, 3.3, 0.07, palette["accent"], radius=False)
    add_text(slide, meta.get("title", "G5 SGD Shell"), 0.30, 5.24, 7.8, 0.18, size=7, color="555555")
    add_text(slide, f"{slide_no}/{slide_count}", 8.9, 5.21, 0.75, 0.22, size=8, color="555555", align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, label, text, palette, body_size=13):
    add_rect(slide, x, y, w, h, palette["softBg"], radius=True)
    if label:
        add_text(slide, label, x+0.14, y+0.08, w-0.28, 0.30, size=13, color=palette["banner"], bold=True)
        add_text(slide, text, x+0.14, y+0.43, w-0.28, h-0.51, size=body_size, color="202124", valign=MSO_ANCHOR.TOP)
    else:
        add_text(slide, text, x+0.14, y+0.12, w-0.28, h-0.24, size=body_size, color="202124", valign=MSO_ANCHOR.TOP)


def render_content(slide, slide_data, palette, index):
    content = slide_data.get("content", [])
    step = slide_data["step"]
    if step == "cover":
        add_text(slide, slide_data["title"], 0.65, 1.35, 8.7, 1.35, size=34, color=palette["banner"], bold=True, align=PP_ALIGN.CENTER)
        y = 2.95
        for block in content:
            if block.get("type") == "text":
                add_text(slide, block.get("text", ""), 1.2, y, 7.6, 0.55, size=15, color="454545", align=PP_ALIGN.CENTER); y += 0.6
        return
    add_text(slide, slide_data["title"], 0.55, 0.92, 8.9, 0.48, size=23, color=palette["banner"], bold=True)
    if step == "case":
        y = 1.58
        for block in content:
            if block.get("type") == "tag":
                x = 0.72
                for tag in block.get("items", [])[:4]:
                    add_rect(slide, x, y, 1.35, 0.35, palette["banner"], radius=True)
                    add_text(slide, tag, x+0.04, y+0.02, 1.27, 0.29, size=8, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
                    x += 1.5
                y += 0.52
            else:
                add_card(slide, 0.72, y, 8.55, 1.55, block.get("label", "Snapshot"), block.get("text", ""), palette)
                y += 1.72
        return
    if step == "questions" or any(b.get("type") == "numbered-list" for b in content):
        items = []
        for b in content:
            if b.get("type") == "numbered-list": items.extend(b.get("items", []))
            elif b.get("type") == "text": items.append(b.get("text", ""))
        y = 1.52
        for i, item in enumerate(items[:7], 1):
            add_rect(slide, 0.72, y, 0.46, 0.42, palette["banner"], radius=True)
            add_text(slide, str(i), 0.75, y+0.01, 0.40, 0.38, size=12, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
            add_text(slide, item, 1.35, y, 7.8, 0.44, size=14, color="202124")
            y += 0.53
        return
    if step == "visuals":
        table_block = next((b for b in content if b.get("type") == "table"), None)
        if table_block:
            cols = table_block.get("columns", [])
            rows = table_block.get("rows", [])
            x0, y0, width = 0.72, 1.52, 8.55
            col_w = width / max(1, len(cols))
            for ci, col in enumerate(cols):
                add_rect(slide, x0+ci*col_w, y0, col_w-0.03, 0.42, palette["banner"], radius=False)
                add_text(slide, col, x0+ci*col_w+0.05, y0+0.02, col_w-0.13, 0.37, size=11, color="FFFFFF", bold=True)
            for ri, row in enumerate(rows[:5]):
                yy = y0 + 0.44 + ri*0.49
                for ci, val in enumerate(row[:len(cols)]):
                    add_rect(slide, x0+ci*col_w, yy, col_w-0.03, 0.45, palette["softBg"] if ri%2==0 else "F7F7F7", radius=False)
                    add_text(slide, val, x0+ci*col_w+0.05, yy+0.02, col_w-0.13, 0.40, size=10, color="202124")
        flow = next((b for b in content if b.get("type") == "flowchart"), None)
        if flow:
            nodes = flow.get("nodes", [])[:3]
            start_x, y, box_w, gap = 0.75, 4.15, 2.35, 0.88
            for i, node in enumerate(nodes):
                x = start_x + i*(box_w+gap)
                add_rect(slide, x, y, box_w, 0.58, palette["softBg"], radius=True, line=palette["banner"])
                add_text(slide, node.get("label", node.get("id", "Step")), x+0.08, y+0.06, box_w-0.16, 0.46, size=10, color=palette["banner"], bold=True, align=PP_ALIGN.CENTER)
            for i in range(min(len(nodes)-1, 2)):
                x = start_x + (i+1)*box_w + i*gap + 0.08
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y+0.18), Inches(gap-0.16), Inches(0.22))
                arrow.fill.solid(); arrow.fill.fore_color.rgb = rgb(palette["accent"]); arrow.line.fill.background()
        return
    if step in ("answers", "takeaways", "references", "thank-you"):
        blocks = []
        for b in content:
            if b.get("type") == "card": blocks.append((b.get("label", ""), b.get("text", "")))
            elif b.get("type") == "text": blocks.append((b.get("label", ""), b.get("text", "")))
            elif b.get("type") == "citation": blocks.append((b.get("label", ""), b.get("text", "")))
            elif b.get("type") == "numbered-list": blocks.extend((f"{i}.", t) for i,t in enumerate(b.get("items", []), 1))
        if step == "thank-you":
            y = 2.0
            for label, text in blocks:
                add_text(slide, text or label, 0.7, y, 8.6, 0.42, size=17 if y==2.0 else 11, color=palette["banner"], bold=(y==2.0), align=PP_ALIGN.CENTER); y += 0.55
            if not blocks: add_text(slide, "GROUP 5", 0.7, 2.3, 8.6, 0.6, size=25, color=palette["banner"], bold=True, align=PP_ALIGN.CENTER)
            return
        if not blocks: blocks = [("Content", "Add reviewed content to this slide.")]
        if len(blocks) == 1:
            add_card(slide, 0.72, 1.58, 8.55, 3.1, blocks[0][0], blocks[0][1], palette)
        else:
            card_w = 4.16
            for i, (label, text) in enumerate(blocks[:4]):
                x = 0.72 + (i%2)*4.39
                y = 1.58 + (i//2)*1.72
                add_card(slide, x, y, card_w, 1.48, label, text, palette, body_size=11)
        return


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("outline", type=Path, nargs="?", default=ROOT / "examples" / "demo-outline.json")
    parser.add_argument("--output", type=Path, default=ROOT / "dist" / "g5-demo-deck.pptx")
    args = parser.parse_args()
    theme = json.loads(THEME_PATH.read_text(encoding="utf-8"))
    data = json.loads(args.outline.read_text(encoding="utf-8"))
    errors = validate_outline(data, theme)
    if errors:
        print("Outline validation failed:", file=sys.stderr)
        for error in errors: print(f"- {error}", file=sys.stderr)
        return 2
    meta = data["metadata"]
    meta.setdefault("roster", theme["branding"]["defaultRoster"])
    prs = Presentation()
    prs.slide_width = Inches(10); prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]
    palette = theme["subjects"][meta["courseCode"]]
    count = len(data["slides"])
    for idx, slide_data in enumerate(data["slides"], 1):
        slide = prs.slides.add_slide(blank)
        add_shell(slide, meta, theme, idx, count)
        render_content(slide, slide_data, palette, idx)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    print(f"Built {args.output} ({len(prs.slides)} slides, {meta['courseCode']}, editable shapes/text).")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
