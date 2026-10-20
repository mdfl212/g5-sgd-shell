#!/usr/bin/env python3
"""Inspect the supplied PowerPoint shell and emit a machine-readable report."""
from __future__ import annotations
import argparse, json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PPTX = ROOT / "assets" / "G5_SGD_Shell_Template_MD201-MD207.pptx"
DEFAULT_OUT = ROOT / "docs" / "template-inspection.json"
EMU_PER_INCH = 914400


def shape_record(shape):
    rec = {"name": shape.name, "type": str(shape.shape_type),
           "leftIn": round(shape.left / EMU_PER_INCH, 3),
           "topIn": round(shape.top / EMU_PER_INCH, 3),
           "widthIn": round(shape.width / EMU_PER_INCH, 3),
           "heightIn": round(shape.height / EMU_PER_INCH, 3)}
    if getattr(shape, "has_text_frame", False):
        rec["text"] = shape.text.strip().replace("\n", " / ")[:300]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE: rec["image"] = True
    return rec


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", type=Path, default=DEFAULT_PPTX)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    if not args.pptx.is_file(): parser.error(f"PowerPoint file not found: {args.pptx}")
    prs = Presentation(str(args.pptx))
    report = {
        "file": str(args.pptx),
        "slideCount": len(prs.slides),
        "dimensionsInches": {"width": round(prs.slide_width / EMU_PER_INCH, 3),
                             "height": round(prs.slide_height / EMU_PER_INCH, 3),
                             "aspectRatio": round(prs.slide_width / prs.slide_height, 4)},
        "masterCount": len(prs.slide_masters),
        "layouts": [{"index": i, "name": layout.name, "placeholderCount": len(layout.placeholders)}
                    for i, layout in enumerate(prs.slide_layouts)],
        "slides": [],
        "limitations": ["Structural inventory only; not a clinical review.",
                        "Text fit and visual quality require rendering and visual inspection."]
    }
    for index, slide in enumerate(prs.slides, start=1):
        text_shapes = [s for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text.strip()]
        report["slides"].append({
            "number": index, "layout": slide.slide_layout.name, "shapeCount": len(slide.shapes),
            "textShapeCount": len(text_shapes),
            "imageCount": sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE),
            "texts": [s.text.strip().replace("\n", " / ")[:180] for s in text_shapes[:20]],
            "shapes": [shape_record(s) for s in slide.shapes]})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {args.output}")
    print(f"Slides: {report['slideCount']} | size: {report['dimensionsInches']['width']} × {report['dimensionsInches']['height']} in | layouts: {len(report['layouts'])} | masters: {len(prs.slide_masters)}")

if __name__ == "__main__":
    main()
