#!/usr/bin/env python3
"""Validate basic structural requirements for a generated G5 deck."""
import argparse, json
from pathlib import Path
from pptx import Presentation
EXPECTED = ["cover", "case", "questions", "answers", "visuals", "takeaways", "references", "thank-you"]


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("pptx", type=Path)
    p.add_argument("--outline", type=Path, help="Optional source outline to check spine against")
    args = p.parse_args()
    errors = []
    try:
        prs = Presentation(str(args.pptx))
    except Exception as e:
        print(f"FAIL: cannot open PowerPoint: {e}")
        return 1
    if len(prs.slides) < 8:
        errors.append(f"Expected at least 8 slides; found {len(prs.slides)}")
    w, h = prs.slide_width/914400, prs.slide_height/914400
    if abs(w-10)>0.02 or abs(h-5.625)>0.02:
        errors.append(f"Unexpected dimensions: {w:.3f} × {h:.3f} in")
    if args.outline:
        data = json.loads(args.outline.read_text(encoding="utf-8"))
        steps = []
        for s in data.get("slides", []):
            if s.get("step") not in steps:
                steps.append(s.get("step"))
        if steps != EXPECTED:
            errors.append("Outline step order invalid: " + " -> ".join(str(x) for x in steps))
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.left < -1000 or sh.top < -1000 or sh.left+sh.width > prs.slide_width+1500 or sh.top+sh.height > prs.slide_height+1500:
                errors.append(f"Slide {si}: shape '{sh.name}' extends beyond slide bounds")
    if errors:
        print("G5 deck validation: FAIL")
        for e in errors: print("-", e)
        return 1
    print(f"G5 deck validation: PASS ({len(prs.slides)} slides, {w:.3f} × {h:.3f} in)")
    print("Checks: PowerPoint opens, minimum slide count, 16:9 dimensions, basic shape bounds.")
    print("Not checked: text overflow, visual quality, citation truth, clinical correctness, or source fidelity.")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
